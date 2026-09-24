"""
export.py — Esportazione dell'inventario dispositivi (Excel) e del report scadenze (PDF)
"""

from datetime import datetime
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

COLORE_INTESTAZIONE = "2E5090"


def esporta_dispositivi_excel(dispositivi: List[dict], path: str) -> None:
    """Esporta l'elenco dei dispositivi (già filtrato, se applicabile) in un foglio Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Dispositivi"

    intestazioni = [
        "ID", "Nome", "Categoria", "Produttore", "Modello", "N. Seriale",
        "UDI", "Stato", "Costo (€)", "Scadenza Garanzia",
    ]
    ws.append(intestazioni)

    font_intestazione = Font(name="Arial", bold=True, color="FFFFFF")
    riempimento_intestazione = PatternFill(start_color=COLORE_INTESTAZIONE, end_color=COLORE_INTESTAZIONE, fill_type="solid")
    for col in range(1, len(intestazioni) + 1):
        cella = ws.cell(row=1, column=col)
        cella.font = font_intestazione
        cella.fill = riempimento_intestazione
        cella.alignment = Alignment(horizontal="center")

    for d in dispositivi:
        ws.append([
            d["id"], d["nome"], d["categoria"], d["produttore"], d.get("modello") or "",
            d["numero_seriale"], d.get("udi") or "", d["stato"],
            d.get("costo") or "", d.get("data_scadenza_garanzia") or "",
        ])
        for cella in ws[ws.max_row]:
            cella.font = Font(name="Arial")

    for col in range(1, len(intestazioni) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    ws.freeze_panes = "A2"
    wb.save(path)


def esporta_scadenze_pdf(scadenze: List[dict], entro_giorni: int, path: str) -> None:
    """Esporta il report delle scadenze (manutenzioni/calibrazioni) in PDF, con le voci scadute evidenziate."""
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    stili = getSampleStyleSheet()
    storia = []

    storia.append(Paragraph("Report Scadenze — Dispositivi Medici", stili["Title"]))
    storia.append(Paragraph(
        f"Generato il {datetime.now().strftime('%d/%m/%Y %H:%M')} · Finestra: prossimi {entro_giorni} giorni",
        stili["Normal"],
    ))
    storia.append(Spacer(1, 16))

    if scadenze:
        dati_tabella = [["Dispositivo", "Tipo", "Scadenza", "Giorni rimanenti"]]
        for s in scadenze:
            giorni = s["giorni_rimanenti"]
            etichetta = f"SCADUTO da {abs(giorni)} gg" if giorni < 0 else f"{giorni} gg"
            dati_tabella.append([s["dispositivo_nome"], s["tipo_scadenza"], s["data_scadenza"], etichetta])

        tabella = Table(dati_tabella, colWidths=[6 * cm, 3 * cm, 3 * cm, 4 * cm])
        stile_tabella = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{COLORE_INTESTAZIONE}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ])
        for i, s in enumerate(scadenze, start=1):
            if s["giorni_rimanenti"] < 0:
                stile_tabella.add("TEXTCOLOR", (3, i), (3, i), colors.red)
        tabella.setStyle(stile_tabella)
        storia.append(tabella)
    else:
        storia.append(Paragraph("Nessuna scadenza nella finestra selezionata.", stili["Normal"]))

    doc.build(storia)
